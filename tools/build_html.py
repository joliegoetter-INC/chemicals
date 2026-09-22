#!/usr/bin/env python3
"""Chemical_Inventory.xlsx -> inventory.html (searchable, offline). Usage: build_html.py [root]"""
import sys, os, json, glob, datetime
from openpyxl import load_workbook
from PIL import Image, ImageOps
root=os.path.abspath(sys.argv[1] if len(sys.argv)>1 else os.path.join(os.path.dirname(__file__),".."))
wb=load_workbook(os.path.join(root,"Chemical_Inventory.xlsx"))
def rows(tab):
    it=wb[tab].iter_rows(values_only=True); head=next(it)
    return [dict(zip(head,[("" if v is None else v) for v in r])) for r in it]
chem={c["CAS"]:c for c in rows("Chemicals")}
os.makedirs(os.path.join(root,"photos","thumbs"),exist_ok=True)
items=[]
for r in rows("Inventory"):
    ph=r["Photo"]; src=os.path.join(root,"photos",ph) if ph else ""
    thumb=""
    if ph and os.path.exists(src):
        t=os.path.join(root,"photos","thumbs",ph)
        if not os.path.exists(t):
            im=ImageOps.exif_transpose(Image.open(src)); im.thumbnail((420,560)); im.convert("RGB").save(t,quality=80)
        thumb="photos/thumbs/"+ph
    c=chem.get(r["CAS"],{})
    items.append({"id":r["Bottle ID"],"name":r["Name"],"nick":r["Common name"],"cas":r["CAS"],"vendor":r["Vendor"],"cat":r["Catalog no."],
      "size":f'{r["Size"]} {r["Unit"]}'.strip(),"state":r["Physical state"],"purity":r["Purity/grade"],"lot":r["Lot"],"loc":r["Location"],
      "flag":r["Flag"],"why":r["Flag reason"],"signal":c.get("Signal word") or r["Signal word (label)"],"codes":c.get("Hazard codes",""),
      "haz":c.get("Hazard summary",""),"received":r["Received (YYYY-MM)"],"receivedRaw":r["Received as written"],"exp":r["Expiration"],
      "storage":r["Storage note"],"photo":("photos/"+ph) if thumb else "","thumb":thumb,"source":r["Source"],"conf":r["Read confidence"],
      "notes":r["Read notes"],"recheck":r["Recheck"]=="YES","recheckWhy":r["Recheck reason"],"sds":r["SDS file"],"pubchem":c.get("Name (PubChem)",""),"sizeN":r["Size"],"unit":r["Unit"],"opened":r["Opened"],"edited":r["Edited by hand"],"catg":r["Category"],"sub":r["Subcategory"],"zone":r["Proposed location"]})
html=open(os.path.join(os.path.dirname(os.path.abspath(__file__)),"inventory_template.html")).read()
html=html.replace("/*DATA*/[]",json.dumps(items,ensure_ascii=False).replace("</","<\\/")).replace("{{BUILT}}",datetime.date.today().isoformat())
open(os.path.join(root,"inventory.html"),"w").write(html)
print(len(items),"items ->",os.path.join(root,"inventory.html"))
