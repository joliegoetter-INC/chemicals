#!/usr/bin/env python3
"""Merge label readings (JSON) -> PubChem hazard lookup -> Chemical_Inventory.xlsx
Usage: build_inventory.py readings_dir out.xlsx"""
import sys, json, glob, re, os, time, urllib.request, urllib.parse, datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.formatting.rule import FormulaRule

REPRO = {"H360","H361","H362","H340","H341"}
# Common peroxide formers (CAS). Not exhaustive: extend as needed.
PEROXIDE = {"60-29-7":"diethyl ether","109-99-9":"THF","123-91-1":"1,4-dioxane","108-20-3":"diisopropyl ether",
 "110-71-4":"glyme","111-96-6":"diglyme","112-49-2":"triglyme","143-24-8":"tetraglyme","96-47-9":"2-MeTHF",
 "1634-04-4":"MTBE","100-42-5":"styrene","108-05-4":"vinyl acetate",
 "110-83-8":"cyclohexene","98-82-8":"cumene","67-63-0":"isopropanol","78-92-2":"2-butanol","119-64-2":"tetralin",
 "91-17-8":"decalin","75-35-4":"vinylidene chloride","106-99-0":"butadiene","79-10-7":"acrylic acid",
 "80-62-6":"methyl methacrylate","107-13-1":"acrylonitrile","5614-37-9":"cyclopentyl methyl ether","142-96-1":"dibutyl ether",
 "629-14-1":"ethylene glycol diethyl ether","100-51-6":"benzyl alcohol","108-10-1":"MIBK","75-07-0":"acetaldehyde","98-01-1":"furfural"}

def cas_ok(c):
    m=re.fullmatch(r"(\d{2,7})-(\d\d)-(\d)",c or "")
    if not m: return False
    d=(m.group(1)+m.group(2))[::-1]
    return sum(int(x)*(i+1) for i,x in enumerate(d))%10==int(m.group(3))

def get(url):
    for a in range(3):
        try:
            with urllib.request.urlopen(urllib.request.Request(url,headers={"User-Agent":"chem-inventory"}),timeout=30) as r: return r.read().decode()
        except Exception as e:
            if "404" in str(e): return None
            time.sleep(1.5)
    return None

def pubchem(cas, cache):
    if cas in cache: return cache[cas]
    out={"found":False,"codes":{},"signal":"","name":"","formula":"","mw":""}
    t=get(f"https://pubchem.ncbi.nlm.nih.gov/rest/pug/compound/name/{cas}/cids/TXT")
    if t and t.strip().split():
        cids=t.split()[:4]; cid=cids[0]; out["found"]=True
        for c in cids:
            g0=get(f"https://pubchem.ncbi.nlm.nih.gov/rest/pug_view/data/compound/{c}/JSON?heading=GHS+Classification")
            if g0 and re.search(r'"H\d{3}',g0): cid=c; break
        out["cid"]=cid
        p=get(f"https://pubchem.ncbi.nlm.nih.gov/rest/pug/compound/cid/{cid}/property/Title,MolecularFormula,MolecularWeight/JSON")
        if p:
            pr=json.loads(p)["PropertyTable"]["Properties"][0]; out.update(name=pr.get("Title",""),formula=pr.get("MolecularFormula",""),mw=pr.get("MolecularWeight",""))
        g=get(f"https://pubchem.ncbi.nlm.nih.gov/rest/pug_view/data/compound/{cid}/JSON?heading=GHS+Classification")
        if g:
            for code,pct,txt in re.findall(r'"(H\d{3}[A-Za-z]*)(?: \(([\d.]+)%\))?: ([^"\[]+)',g):
                pct=float(pct) if pct else -1.0
                if code not in out["codes"] or pct>out["codes"][code][0]: out["codes"][code]=(pct,txt.strip())
            try:
                d=json.loads(g); refs={x["ReferenceNumber"]:x["SourceName"] for x in d["Record"]["Reference"]}; src={}
                def walk(n):
                    if isinstance(n,dict):
                        if "ReferenceNumber" in n and "Value" in n:
                            for sm in n["Value"].get("StringWithMarkup",[]):
                                mm=re.match(r"(H\d{3}[A-Za-z]*)",sm["String"])
                                if mm: src.setdefault(mm.group(1),[]).append(refs.get(n["ReferenceNumber"],"?"))
                        for v in n.values(): walk(v)
                    elif isinstance(n,list):
                        for v in n: walk(v)
                walk(d); out["src"]={k:sorted(set(v)) for k,v in src.items()}
            except Exception: out["src"]={}
            m=re.search(r'"String": "(Danger|Warning)"',g); out["signal"]=m.group(1) if m else ""
        time.sleep(0.3)
    cache[cas]=out; return out

def vendor(v):
    v=v or ""
    for k,n in (("TCI","TCI"),("Tokyo Chemical","TCI"),("Sartomer","Sartomer (Arkema)"),("Arkema","Sartomer (Arkema)"),("Sigma","Sigma-Aldrich"),("AmBeed","AmBeed"),("AGC","AGC Si-Tech"),("NAGASE","Nagase")):
        if k.lower() in v.lower(): return n
    return v.strip()

def cas_from_name(name, cache):
    """PubChem exact-name lookup. Returns a CAS only if the name resolves and a checksum-valid CAS is in its synonyms."""
    key="name:"+name
    if key in cache: return cache[key]
    cas=""
    t=get("https://pubchem.ncbi.nlm.nih.gov/rest/pug/compound/name/"+urllib.parse.quote(name)+"/synonyms/JSON")
    if t:
        try:
            syn=json.loads(t)["InformationList"]["Information"][0]["Synonym"]
            cas=next((x for x in syn if cas_ok(x)),"")
        except Exception: pass
    cache[key]=cas; return cas

def sds_index(root):
    idx={}
    for f in glob.glob(os.path.join(root,"sds","*.pdf")):
        parts=os.path.basename(f)[:-4].split("_")
        if len(parts)>=2: idx[parts[1].lower()]="sds/"+os.path.basename(f)
    for rf in glob.glob(os.path.join(root,"sds","results_*.json")):
        try:
            for x in json.load(open(rf)):
                if isinstance(x,dict) and x.get("status")=="downloaded" and x.get("file"):
                    for k in (x.get("catalog_no"),x.get("product")):
                        if k: idx[str(k).lower()]="sds/"+os.path.basename(x["file"]); idx[str(k).split("-")[0].lower()]="sds/"+os.path.basename(x["file"])
        except Exception: pass
    return idx

CAT_ORDER=["TPU & polymers","Additives","Monomers & oligomers","Uncategorized"]
def categorize(r,cas,cats):
    if cas in cats: return cats[cas]
    n=((r.get("name") or "")+" "+(r.get("common_name") or "")).lower()
    for k,v in cats.items():
        if not k.startswith("_") and not re.match(r"\d",k) and k in n: return v
    return ["Uncategorized",""]

def zone(cat,sub,flag):
    """Proposed storage zone. RED overrides everything: one Carl-only shelf."""
    if flag=="RED": return "Z7 Carl-only (red dot)"
    if "drum" in sub: return "FLOOR Corrosives spill tray"
    if cat=="TPU & polymers": return "Z1 TPU & polymers (dry shelf)"
    if sub in("Dye","Photoinitiator"): return "Z3 Light-sensitive: "+("dyes bin" if sub=="Dye" else "photoinitiators bin")
    if cat=="Additives": return "Z4 Additives: liquids" if sub in("Plasticizer","Plasticizer / humectant","Solvent","Amine") else "Z2 Additives: powders"
    if sub.startswith("Oligomer"): return "Z6 Oligomers & resins (cabinet, bottom)"
    if cat=="Monomers & oligomers": return "Z5 Monomers (flammables cabinet): "+("methacrylates shelf" if "methacryl" in sub else "acrylates shelf")
    return ""

EDITABLE={"Name":"name","Common name":"common_name","CAS":"cas","Vendor":"vendor","Catalog no.":"catalog_no","Size":"size","Unit":"unit",
 "Physical state":"physical_state","Purity/grade":"purity","Lot":"lot","Location":"location","Received (YYYY-MM)":"received","Opened":"opened",
 "Expiration":"expiration","Storage note":"storage_note","Category":"category","Subcategory":"subcategory","Read notes":"problems"}
def norm(v): 
    v="" if v is None else v
    if isinstance(v,float) and v==int(v): v=int(v)
    return str(v).strip()

def pull_excel_edits(src,dst):
    """Excel is the editing surface: anything a person changed in an editable column since the last build
    is saved to edits.json, and new rows without a Bottle ID become manual entries. Runs before every rebuild."""
    snap_p=os.path.join(src,"last_build.json")
    if not (os.path.exists(dst) and os.path.exists(snap_p)): return 0
    from openpyxl import load_workbook
    snap=json.load(open(snap_p)); ws=load_workbook(dst,data_only=True)["Inventory"]; it=ws.iter_rows(values_only=True); head=list(next(it))
    ep=os.path.join(src,"edits.json"); edits=json.load(open(ep)) if os.path.exists(ep) else {}
    mp=os.path.join(src,"manual.json"); manual=json.load(open(mp)) if os.path.exists(mp) else []
    today=datetime.date.today().isoformat(); n=0
    for row in it:
        d={h:norm(v) for h,v in zip(head,row) if h}
        bid=d.get("Bottle ID","")
        if not bid:
            if d.get("Name") or d.get("Common name"):
                m={k:d.get(h,"") for h,k in EDITABLE.items() if d.get(h)}
                m.update(key="manual-"+datetime.datetime.now().strftime("%Y%m%d-%H%M%S")+f"-{len(manual)}",source="manual",confidence="",t_start=10**6+len(manual),_edited=today)
                manual.append(m); n+=1
            continue
        old=snap.get(bid)
        if not old: continue
        ch={k:d.get(h,"") for h,k in EDITABLE.items() if d.get(h,"")!=old.get(h,"")}
        if old.get("Recheck")=="YES" and d.get("Recheck","").upper()!="YES": ch["recheck_done"]=True
        if ch: e=edits.get(bid,{}); e.update(ch); e["_edited"]=today; edits[bid]=e; n+=1
    if n:
        json.dump(edits,open(ep,"w"),indent=1,ensure_ascii=False); json.dump(manual,open(mp,"w"),indent=1,ensure_ascii=False)
    return n

def main(src,dst):
    n_edits=pull_excel_edits(src,dst)
    if n_edits: print(n_edits,"change(s) picked up from the Excel file")
    rows=[]
    for f in sorted(glob.glob(os.path.join(src,"seg_*.json"))): rows+=json.load(open(f))
    rows.sort(key=lambda r:r.get("t_start",0))
    # de-dupe boundary overlaps: same cas+lot+catalog within 8s
    ded=[]
    for r in rows:
        k=(r.get("cas"),r.get("lot"),r.get("catalog_no"),r.get("received"),r.get("common_name"))
        if ded and (k[0] or k[1]) and k==ded[-1][0] and abs(r.get("t_start",0)-ded[-1][1].get("t_end",0))<8 : continue
        ded.append((k,r))
    rows=[r for _,r in ded]
    # manual entries (added from the page) come after the scanned ones
    mp=os.path.join(src,"manual.json")
    if os.path.exists(mp): rows+=json.load(open(mp))
    # stable bottle IDs: once a bottle has an ID it keeps it, new bottles get the next number
    ip=os.path.join(src,"ids.json"); ids=json.load(open(ip)) if os.path.exists(ip) else {}
    def rkey(r): return r.get("key") or os.path.basename(r.get("photo") or "") or f"{r.get('name')}|{r.get('lot')}|{r.get('t_start')}"
    nxt=max([int(v[1:]) for v in ids.values()]+[0])
    for r in rows:
        k=rkey(r)
        if k not in ids: nxt+=1; ids[k]=f"B{nxt:04d}"
        r["_id"]=ids[k]
    json.dump(ids,open(ip,"w"),indent=1)
    # human edits override what was read from the label
    ep=os.path.join(src,"edits.json"); edits=json.load(open(ep)) if os.path.exists(ep) else {}
    for r in rows:
        e=edits.get(r["_id"],{})
        for k,v in e.items():
            if k!="_edited": r[k]=v
        if e: r["_edited"]=e.get("_edited","")
    cpath=os.path.join(src,"pubchem_cache.json"); cache=json.load(open(cpath)) if os.path.exists(cpath) else {}
    today=datetime.date.today().isoformat(); root=os.path.dirname(os.path.abspath(dst)); sidx=sds_index(root); cats=json.load(open(os.path.join(src,'categories.json'))) if os.path.exists(os.path.join(src,'categories.json')) else {}
    inv,chem,review,carl=[],{},[],[]
    for i,r in enumerate(rows,1):
        bid=r["_id"]; cas=(r.get("cas") or "").strip(); flag="";why=[];probs=[r.get("problems","")] if r.get("problems") else []
        if cas and not cas_ok(cas): probs.append(f"CAS {cas} fails checksum"); cas=""
        r["vendor"]=vendor(r.get("vendor"))
        if not cas and r.get("name") and r.get("confidence")!="low":
            cas=cas_from_name(re.sub(r"\s*\(.*$","",r["name"]).strip(),cache)
            if cas: why.append("CAS not on label: found by name lookup, confirm against SDS")
        twins=[x["_id"] for x in rows if x is not r and x.get("lot") and x.get("lot")==r.get("lot")]
        recheck=[]
        if twins: recheck.append("Possible duplicate of "+", ".join(twins)+" (same lot): confirm these are separate bottles")
        if r.get("recheck_done"): recheck=[]
        elif r.get("confidence") in ("medium","low"): recheck.append("Label read with "+r["confidence"]+" confidence: re-photograph or check by eye")
        cat=(r.get("catalog_no") or ""); sds=sidx.get(cat.lower()) or sidx.get(cat.split("-")[0].lower()) or sidx.get((r.get("name") or "").lower()) or sidx.get(re.sub(r"[^a-z0-9]","",(r.get("name") or "").lower())) or ""
        if cas:
            pc=pubchem(cas,cache)
            if not pc["found"]: flag="REVIEW"; why.append("CAS not in PubChem (common for polymers): confirm with vendor SDS")
            else:
                codes=pc["codes"]; base={c[:4] for c in codes}
                rep=sorted(c for c in codes if c[:4] in REPRO)
                if rep: flag="RED"; why.append("Reproductive/genetic: "+", ".join(c+(f" ({codes[c][0]:.0f}% of company reports)" if codes[c][0]>=0 else "")+" [source: "+"/".join(pc.get("src",{}).get(c,["?"]))+"]" for c in rep))
                if cas in PEROXIDE:
                    why.append("Peroxide former ("+PEROXIDE[cas]+"): track received/opened dates")
                    if flag!="RED": flag="ORANGE"
                    else: flag="RED"
                if not codes: why.append("No GHS data on PubChem (common for polymers and dyes): confirm with vendor SDS"); flag=flag or "REVIEW"
                if cas not in chem:
                    chem[cas]=[cas,r.get("name",""),pc["name"],pc["formula"],pc["mw"],pc["signal"],", ".join(sorted(codes)),
                        "; ".join(f"{c}: {codes[c][1]}" for c in sorted(codes)),flag if flag in("RED","ORANGE") else ("REVIEW" if not codes else "None"),
                        " | ".join(why),("ORANGE too" if flag=="RED" and cas in PEROXIDE else ""),"","PubChem CID "+pc.get("cid",""),today]
        else:
            flag="REVIEW"; why.append("No CAS on label (mixture or trade product?): hazards need the vendor SDS")
            carl.append([r.get("vendor",""),r.get("catalog_no",""),"",r.get("name",""),"No CAS, need vendor SDS","Open",""])
        if r.get("confidence") in ("low",) and flag not in("RED",): flag=flag or "REVIEW"
        mine=[p for p in probs if p and p!=r.get("problems")]
        if r.get("confidence")=="low": mine.append("Label barely readable in video: re-photograph")
        if flag=="REVIEW" or mine: review.append([bid,os.path.basename(r.get("photo","")),"; ".join(mine+([] if mine and flag!="REVIEW" else why)),r.get("name","") or r.get("common_name",""),"",""])
        inv.append([bid,r.get("name",""),r.get("common_name",""),cas,r.get("vendor",""),r.get("catalog_no",""),r.get("size",""),r.get("unit",""),
            r.get("physical_state",""),r.get("purity",""),r.get("lot",""),r.get("location",""),flag," | ".join(why),r.get("signal_word",""),", ".join(r.get("pictograms",[]) or []),
            r.get("received",""),r.get("received_raw",""),r.get("opened",""),r.get("expiration",""),r.get("storage_note",""),os.path.basename(r.get("photo","")),
            ("manual entry" if r.get("source")=="manual" else "photo" if r.get("source")=="photo" else f"video {r.get('t_start','')}-{r.get('t_end','')}s"),r.get("confidence",""),r.get("problems",""),"YES" if recheck else ""," | ".join(recheck),sds]+(lambda c:[c[0],c[1],zone(c[0],c[1],flag)])([r["category"],r.get("subcategory","")] if r.get("category") else categorize(r,cas,cats))+[r.get("_edited","")])
    json.dump(cache,open(cpath,"w"))
    wb=Workbook()
    def tab(ws,title,cols,data,flagcol=None):
        ws.title=title; ws.append(cols)
        for d in data: ws.append(d)
        for c in ws[1]: c.font=Font(bold=True); c.fill=PatternFill("solid",fgColor="DDDDDD")
        ws.freeze_panes="B2"; ws.auto_filter.ref=ws.dimensions
        for i,c in enumerate(cols,1):
            L=ws.cell(1,i).column_letter; w=max([len(str(c))]+[len(str(d[i-1])) for d in data if i-1<len(d)]+[8]); ws.column_dimensions[L].width=min(w+2,45)
        if flagcol:
            last=ws.cell(1,len(cols)).column_letter; rng=f"A2:{last}2000"
            for val,col in (("RED","F4C7C3"),("ORANGE","FCE8B2"),("REVIEW","E0E0E0")):
                ws.conditional_formatting.add(rng,FormulaRule(formula=[f'${flagcol}2="{val}"'],fill=PatternFill("solid",bgColor=col)))
    tab(wb.active,"Inventory",["Bottle ID","Name","Common name","CAS","Vendor","Catalog no.","Size","Unit","Physical state","Purity/grade","Lot","Location","Flag","Flag reason","Signal word (label)","Pictograms (label)","Received (YYYY-MM)","Received as written","Opened","Expiration","Storage note","Photo","Source","Read confidence","Read notes","Recheck","Recheck reason","SDS file","Category","Subcategory","Proposed location","Edited by hand"],sorted(inv,key=lambda x:(CAT_ORDER.index(x[28]) if x[28] in CAT_ORDER else 9,x[29],x[1] or "")),"M")
    tab(wb.create_sheet(),"Chemicals",["CAS","Name (label)","Name (PubChem)","Formula","MW","Signal word","Hazard codes","Hazard summary","Flag","Flag reason","Also","SDS file","Hazard source","Checked"],list(chem.values()),"I")
    tab(wb.create_sheet(),"Needs review",["Bottle ID","Photo","Problem","Name as read","Resolved by","Resolved date"],review)
    tab(wb.create_sheet(),"Carl SDS list",["Vendor","Catalog no.","CAS","Name","Why","Status","SDS file"],carl)
    k=wb.create_sheet(); tab(k,"Key",["Flag","Meaning","Sticker"],[["RED","Reproductive or genetic hazard (H360/H361/H362/H340/H341). Carl handles.","Red dot"],["ORANGE","Time-sensitive (peroxide former). Received and opened dates required.","Orange dot"],["REVIEW","Could not confirm hazards. NOT cleared as safe.","None until resolved"],["(blank)","No reproductive or time-sensitive flag found. Still read the hazard summary.","None"],
        ["",""],["HOW TO EDIT","Green column headers on the Inventory tab can be edited right here in Excel. Save, close Excel, then double-click Rebuild.command (or ask Claude to rebuild). Your changes are kept forever.",""],
        ["ADD A BOTTLE","Type a new row at the bottom of Inventory and leave Bottle ID empty. It gets an ID and hazard flags on the next rebuild.",""],
        ["RECHECK","After checking a bottle by eye, change its Recheck cell from YES to OK.",""],
        ["DO NOT EDIT","Grey column headers (Flag, Flag reason, SDS, Proposed location...) are recalculated on every rebuild. Changes there are thrown away.",""]])
    INV_COLS=[c.value for c in wb["Inventory"][1]]
    json.dump({r[0]:{h:norm(v) for h,v in zip(INV_COLS,r)} for r in inv},open(os.path.join(src,"last_build.json"),"w"),ensure_ascii=False)
    for c in wb["Inventory"][1]:
        if c.value in EDITABLE or c.value=="Recheck": c.fill=PatternFill("solid",fgColor="C6E0B4")
    wb.save(dst)
    from collections import Counter
    print(len(inv),"bottles;",len(chem),"unique CAS; flags:",dict(Counter(x[12] or "none" for x in inv)))
main(sys.argv[1],sys.argv[2])
